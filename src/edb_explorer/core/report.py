"""Database reports in html / md / pdf / docx / xlsx / txt / json.

``build_report`` gathers everything into a plain dict (JSON-serialisable) and
the ``render_*`` functions turn that dict into each output format, so the GUI,
CLI and MCP server produce identical reports.
"""

from __future__ import annotations

import html
import json
import os
from collections.abc import Callable
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from edb_explorer import __app_name__, __version__
from edb_explorer.core.database import EdbDatabase
from edb_explorer.core.exceptions import ExportError
from edb_explorer.core.values import display_value

REPORT_FORMATS: tuple[str, ...] = ("html", "pdf", "docx", "md", "xlsx", "txt", "json")
REPORT_FORMAT_LABELS = {
    "html": "HTML page (self-contained)",
    "pdf": "PDF document",
    "docx": "Word document (.docx)",
    "md": "Markdown",
    "xlsx": "Excel workbook (.xlsx)",
    "txt": "Plain text",
    "json": "JSON",
}
ProgressCallback = Callable[[str], bool]

SAMPLE_MAX_COLUMNS = 10
SAMPLE_CELL_LIMIT = 80


@dataclass(slots=True)
class ReportOptions:
    title: str | None = None
    case_id: str | None = None
    analyst: str | None = None
    notes: str | None = None
    include_schema: bool = True
    include_system: bool = False
    count_records: bool = True
    compute_hash: bool = True
    sample_rows: int = 5
    tables: list[str] | None = None  # restrict to these tables (all databases)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


# --------------------------------------------------------------------------- #
# Data gathering
# --------------------------------------------------------------------------- #
def build_report(
    databases: list[EdbDatabase], options: ReportOptions | None = None, progress: ProgressCallback | None = None
) -> dict[str, Any]:
    opts = options or ReportOptions()

    def tick(msg: str) -> None:
        if progress and not progress(msg):
            raise ExportError("Report generation cancelled")

    report: dict[str, Any] = {
        "title": opts.title or f"ESE database report - {', '.join(db.path.name for db in databases)}",
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "generator": f"{__app_name__} {__version__}",
        "case_id": opts.case_id,
        "analyst": opts.analyst,
        "notes": opts.notes,
        "options": opts.to_dict(),
        "databases": [],
    }
    for db in databases:
        tick(f"Collecting {db.path.name}")
        if opts.compute_hash:
            tick(f"Hashing {db.path.name}")
            db.compute_sha256()
        info = db.info.to_dict(include_header=True)
        tables_out: list[dict[str, Any]] = []
        wanted = {db.resolve_table_name(t) for t in opts.tables} if opts.tables else None
        total_records = 0
        for t in db.tables(include_system=opts.include_system):
            if wanted is not None and t.name not in wanted:
                continue
            tick(f"{db.path.name}: {t.display_name}")
            if opts.count_records:
                db.count_records(t.name)
            t = db.table(t.name)
            entry = t.to_dict(include_columns=opts.include_schema)
            total_records += t.record_count or 0
            if opts.sample_rows > 0 and not t.is_system:
                page = db.fetch(t.name, limit=opts.sample_rows, max_length=SAMPLE_CELL_LIMIT)
                cols: list[str] = []
                for row in page.rows:
                    for k in row:
                        if k != "_row" and k not in cols:
                            cols.append(k)
                entry["sample_columns"] = cols[:SAMPLE_MAX_COLUMNS]
                entry["sample_rows"] = [
                    {"_row": r["_row"], **{c: r.get(c) for c in cols[:SAMPLE_MAX_COLUMNS] if r.get(c) is not None}}
                    for r in page.rows
                ]
            tables_out.append(entry)
        report["databases"].append(
            {
                "info": info,
                "profile": {"id": db.profile.id, "name": db.profile.name, "description": db.profile.description},
                "table_count_reported": len(tables_out),
                "total_records": total_records if opts.count_records else None,
                "tables": tables_out,
            }
        )
    return report


# --------------------------------------------------------------------------- #
# Shared helpers
# --------------------------------------------------------------------------- #
def _summary_rows(entry: dict[str, Any]) -> list[tuple[str, str]]:
    info = entry["info"]

    def fmt(v: Any) -> str:
        return "-" if v in (None, "") else str(v)

    return [
        ("File", fmt(info.get("path"))),
        ("Size", f"{info.get('size_bytes', 0):,} bytes"),
        ("SHA-256", fmt(info.get("sha256"))),
        ("Detected type", f"{entry['profile']['name']} - {entry['profile']['description']}"),
        ("Page size", f"{info.get('page_size', 0):,} bytes"),
        ("Format", f"0x{info.get('format_version', 0):X} rev {info.get('format_revision')}"),
        ("Created by format", f"0x{info.get('created_version', 0):X} rev {info.get('created_revision')}"),
        ("Shutdown state", fmt(info.get("state"))),
        ("Windows version", fmt(info.get("windows_version"))),
        ("Created (UTC)", fmt(info.get("created"))),
        ("Last attach (UTC)", fmt(info.get("last_attach"))),
        ("Last detach (UTC)", fmt(info.get("last_detach"))),
        ("Tables", str(info.get("table_count"))),
        ("Total records (reported tables)", fmt(entry.get("total_records"))),
    ]


def _table_rows(entry: dict[str, Any]) -> list[list[str]]:
    out = []
    for t in entry["tables"]:
        out.append(
            [
                t["name"],
                t["display_name"] if t["display_name"] != t["name"] else "",
                str(t["column_count"]),
                str(t["index_count"]),
                "-" if t.get("record_count") is None else f"{t['record_count']:,}",
                t.get("description") or "",
            ]
        )
    return out


def _meta_rows(report: dict[str, Any]) -> list[tuple[str, str]]:
    rows = [("Generated", report["generated_at"]), ("Generator", report["generator"])]
    if report.get("case_id"):
        rows.append(("Case", report["case_id"]))
    if report.get("analyst"):
        rows.append(("Analyst", report["analyst"]))
    return rows


def _cell(v: Any) -> str:
    return display_value(v, None, SAMPLE_CELL_LIMIT)


# --------------------------------------------------------------------------- #
# Renderers
# --------------------------------------------------------------------------- #
def render_markdown(report: dict[str, Any]) -> str:
    L: list[str] = [f"# {report['title']}", ""]
    for k, v in _meta_rows(report):
        L.append(f"- **{k}:** {v}")
    if report.get("notes"):
        L += ["", "## Notes", "", report["notes"]]
    for entry in report["databases"]:
        info = entry["info"]
        L += ["", f"## {info['file_name']}", "", "| Property | Value |", "|---|---|"]
        for k, v in _summary_rows(entry):
            L.append(f"| {k} | {_md(v)} |")
        L += [
            "",
            "### Tables",
            "",
            "| Table | Display name | Columns | Indexes | Records | Description |",
            "|---|---|---:|---:|---:|---|",
        ]
        for r in _table_rows(entry):
            L.append("| " + " | ".join(_md(c) for c in r) + " |")
        for t in entry["tables"]:
            if not (t.get("columns") or t.get("sample_rows")):
                continue
            L += ["", f"### {t['display_name']}" + (f" (`{t['name']}`)" if t["display_name"] != t["name"] else "")]
            if t.get("description"):
                L += ["", t["description"]]
            if t.get("columns"):
                L += ["", "| ID | Column | Type | Storage | Size | Encoding |", "|---:|---|---|---|---:|---|"]
                for c in t["columns"]:
                    L.append(
                        f"| {c['identifier']} | {_md(c['name'])} | {c['type']} | {c['storage']} | {c['size'] or ''} | {c['encoding'] or ''} |"
                    )
                if t.get("indexes"):
                    L += [
                        "",
                        "Indexes: "
                        + ", ".join(
                            f"`{i['name']}`"
                            + (" (primary)" if i["is_primary"] else "")
                            + (f" [{', '.join(i['columns'])}]" if i["columns"] else "")
                            for i in t["indexes"]
                        ),
                    ]
            if t.get("sample_rows"):
                cols = t["sample_columns"]
                L += [
                    "",
                    f"Sample rows ({len(t['sample_rows'])}):",
                    "",
                    "| _row | " + " | ".join(_md(c) for c in cols) + " |",
                    "|---:|" + "---|" * len(cols),
                ]
                for r in t["sample_rows"]:
                    L.append(f"| {r['_row']} | " + " | ".join(_md(_cell(r.get(c))) for c in cols) + " |")
    L.append("")
    return "\n".join(L)


def _md(text: Any) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def render_text(report: dict[str, Any]) -> str:
    L: list[str] = [report["title"], "=" * len(report["title"]), ""]
    for k, v in _meta_rows(report):
        L.append(f"{k:<12} {v}")
    if report.get("notes"):
        L += ["", "Notes", "-----", report["notes"]]
    for entry in report["databases"]:
        info = entry["info"]
        L += ["", "", info["file_name"], "-" * len(info["file_name"]), ""]
        for k, v in _summary_rows(entry):
            L.append(f"  {k:<32} {v}")
        L += ["", "  Tables:", ""]
        rows = _table_rows(entry)
        widths = [
            max(len(h), *(len(r[i]) for r in rows)) if rows else len(h)
            for i, h in enumerate(("Table", "Display name", "Columns", "Indexes", "Records"))
        ]
        head = ["Table", "Display name", "Columns", "Indexes", "Records"]
        L.append("  " + "  ".join(h.ljust(widths[i]) for i, h in enumerate(head)))
        for r in rows:
            L.append("  " + "  ".join(r[i].ljust(widths[i]) for i in range(5)))
        for t in entry["tables"]:
            if not (t.get("columns") or t.get("sample_rows")):
                continue
            L += ["", f"  [{t['display_name']}]" + (f"  ({t['name']})" if t["display_name"] != t["name"] else "")]
            if t.get("description"):
                L.append(f"  {t['description']}")
            for c in t.get("columns") or []:
                L.append(
                    f"    {c['identifier']:>6}  {c['name']:<40} {c['type']:<14} {c['storage']:<9} {c['size'] or ''}"
                )
            for i in t.get("indexes") or []:
                L.append(f"    index {i['name']}{' (primary)' if i['is_primary'] else ''}: {', '.join(i['columns'])}")
            for r in t.get("sample_rows") or []:
                L.append(f"    row {r['_row']}: " + "; ".join(f"{k}={_cell(v)}" for k, v in r.items() if k != "_row"))
    L.append("")
    return "\n".join(L)


_HTML_CSS = """
body{font-family:-apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;margin:2rem auto;max-width:1200px;color:#1f2328;line-height:1.45}
h1{border-bottom:2px solid #2f6fd6;padding-bottom:.3rem}h2{margin-top:2.2rem;color:#1a365d}h3{margin-top:1.6rem}
table{border-collapse:collapse;width:100%;margin:.6rem 0;font-size:13px}th,td{border:1px solid #d0d7de;padding:4px 8px;text-align:left;vertical-align:top}
th{background:#eef2f7}tr:nth-child(even) td{background:#f7f8fa}td.num,th.num{text-align:right}
.meta{color:#6e7781;font-size:13px}.kv td:first-child{font-weight:600;width:16rem;white-space:nowrap}
code{font-family:Consolas,Menlo,monospace;font-size:12px;background:#f0f2f5;padding:1px 4px;border-radius:3px}
.desc{color:#57606a;font-size:13px}.wrap{overflow-x:auto}pre.notes{white-space:pre-wrap;background:#f7f8fa;padding:.8rem;border-radius:6px}
"""


def render_html(report: dict[str, Any]) -> str:
    e = html.escape
    P: list[str] = [
        "<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>",
        f"<title>{e(report['title'])}</title><style>{_HTML_CSS}</style></head><body>",
        f"<h1>{e(report['title'])}</h1><p class='meta'>",
        " &middot; ".join(f"<b>{e(k)}:</b> {e(v)}" for k, v in _meta_rows(report)),
        "</p>",
    ]
    if report.get("notes"):
        P.append(f"<h2>Notes</h2><pre class='notes'>{e(report['notes'])}</pre>")
    if len(report["databases"]) > 1:
        P.append("<h2>Databases</h2><ul>")
        for entry in report["databases"]:
            P.append(
                f"<li><a href='#{e(entry['info']['id'])}'>{e(entry['info']['file_name'])}</a> - {e(entry['profile']['name'])}</li>"
            )
        P.append("</ul>")
    for entry in report["databases"]:
        info = entry["info"]
        P.append(f"<h2 id='{e(info['id'])}'>{e(info['file_name'])}</h2><table class='kv'>")
        for k, v in _summary_rows(entry):
            P.append(f"<tr><td>{e(k)}</td><td>{e(v)}</td></tr>")
        P.append(
            "</table><h3>Tables</h3><div class='wrap'><table><tr><th>Table</th><th>Display name</th>"
            "<th class='num'>Columns</th><th class='num'>Indexes</th><th class='num'>Records</th><th>Description</th></tr>"
        )
        for r in _table_rows(entry):
            P.append(
                f"<tr><td><a href='#{e(info['id'])}-{e(r[0])}'><code>{e(r[0])}</code></a></td><td>{e(r[1])}</td>"
                f"<td class='num'>{r[2]}</td><td class='num'>{r[3]}</td><td class='num'>{r[4]}</td><td class='desc'>{e(r[5])}</td></tr>"
            )
        P.append("</table></div>")
        for t in entry["tables"]:
            if not (t.get("columns") or t.get("sample_rows")):
                continue
            P.append(
                f"<h3 id='{e(info['id'])}-{e(t['name'])}'>{e(t['display_name'])}"
                + (f" <code>{e(t['name'])}</code>" if t["display_name"] != t["name"] else "")
                + "</h3>"
            )
            if t.get("description"):
                P.append(f"<p class='desc'>{e(t['description'])}</p>")
            if t.get("columns"):
                P.append(
                    "<details><summary>Schema: "
                    f"{t['column_count']} columns, {t['index_count']} indexes</summary><div class='wrap'><table>"
                    "<tr><th class='num'>ID</th><th>Column</th><th>Type</th><th>Storage</th><th class='num'>Size</th><th>Encoding</th></tr>"
                )
                for c in t["columns"]:
                    P.append(
                        f"<tr><td class='num'>{c['identifier']}</td><td><code>{e(c['name'])}</code></td><td>{e(c['type'])}</td>"
                        f"<td>{e(c['storage'])}</td><td class='num'>{c['size'] or ''}</td><td>{e(c['encoding'] or '')}</td></tr>"
                    )
                P.append("</table>")
                if t.get("indexes"):
                    P.append(
                        "<p><b>Indexes:</b> "
                        + ", ".join(
                            f"<code>{e(i['name'])}</code>"
                            + (" (primary)" if i["is_primary"] else "")
                            + (f" [{e(', '.join(i['columns']))}]" if i["columns"] else "")
                            for i in t["indexes"]
                        )
                        + "</p>"
                    )
                P.append("</div></details>")
            if t.get("sample_rows"):
                cols = t["sample_columns"]
                P.append(
                    f"<p class='desc'>Sample rows ({len(t['sample_rows'])})</p><div class='wrap'><table><tr><th class='num'>_row</th>"
                    + "".join(f"<th>{e(c)}</th>" for c in cols)
                    + "</tr>"
                )
                for r in t["sample_rows"]:
                    P.append(
                        f"<tr><td class='num'>{r['_row']}</td>"
                        + "".join(f"<td>{e(_cell(r.get(c)))}</td>" for c in cols)
                        + "</tr>"
                    )
                P.append("</table></div>")
    P.append(
        f"<p class='meta'>Generated by {e(report['generator'])}. All database access was read-only.</p></body></html>"
    )
    return "\n".join(P)


def render_pdf(report: dict[str, Any], output: str | os.PathLike[str]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    styles = getSampleStyleSheet()
    body = styles["BodyText"].clone("body", fontSize=8.5, leading=11)
    small = styles["BodyText"].clone("small", fontSize=7, leading=8.5)
    cell = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=report["title"],
        author=report["generator"],
    )
    grid = TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b0b8c0")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dde6f0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7fa")]),
        ]
    )
    kv_style = TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b0b8c0")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2f7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
    )
    esc = html.escape
    W = doc.width
    story: list[Any] = [
        Paragraph(esc(report["title"]), styles["Title"]),
        Paragraph(" &nbsp;&middot;&nbsp; ".join(f"<b>{esc(k)}:</b> {esc(v)}" for k, v in _meta_rows(report)), body),
    ]
    if report.get("notes"):
        story += [
            Spacer(1, 4 * mm),
            Paragraph("Notes", styles["Heading2"]),
            Paragraph(esc(report["notes"]).replace("\n", "<br/>"), body),
        ]
    for n, entry in enumerate(report["databases"]):
        info = entry["info"]
        if n:
            story.append(PageBreak())
        story.append(Paragraph(esc(info["file_name"]), styles["Heading1"]))
        kv = [[Paragraph(esc(k), cell), Paragraph(esc(v), cell)] for k, v in _summary_rows(entry)]
        t = LongTable(kv, colWidths=[45 * mm, W - 45 * mm])
        t.setStyle(kv_style)
        story += [t, Spacer(1, 4 * mm), Paragraph("Tables", styles["Heading2"])]
        data = [[Paragraph(h, cell) for h in ("Table", "Display name", "Cols", "Idx", "Records", "Description")]]
        for r in _table_rows(entry):
            data.append([Paragraph(esc(c), cell) for c in r])
        t = LongTable(data, colWidths=[W * f for f in (0.26, 0.20, 0.06, 0.05, 0.10, 0.33)], repeatRows=1)
        t.setStyle(grid)
        story.append(t)
        for tb in entry["tables"]:
            if not (tb.get("columns") or tb.get("sample_rows")):
                continue
            story += [
                Spacer(1, 3 * mm),
                Paragraph(
                    esc(tb["display_name"])
                    + (f" <font size=7>({esc(tb['name'])})</font>" if tb["display_name"] != tb["name"] else ""),
                    styles["Heading3"],
                ),
            ]
            if tb.get("description"):
                story.append(Paragraph(esc(tb["description"]), small))
            if tb.get("columns"):
                data = [[Paragraph(h, cell) for h in ("ID", "Column", "Type", "Storage", "Size", "Encoding")]]
                for c in tb["columns"]:
                    data.append(
                        [
                            Paragraph(esc(str(x)), cell)
                            for x in (
                                c["identifier"],
                                c["name"],
                                c["type"],
                                c["storage"],
                                c["size"] or "",
                                c["encoding"] or "",
                            )
                        ]
                    )
                t = LongTable(data, colWidths=[W * f for f in (0.08, 0.40, 0.16, 0.12, 0.08, 0.16)], repeatRows=1)
                t.setStyle(grid)
                story.append(t)
                if tb.get("indexes"):
                    story.append(
                        Paragraph(
                            "Indexes: "
                            + ", ".join(
                                esc(i["name"])
                                + (" (primary)" if i["is_primary"] else "")
                                + (f" [{esc(', '.join(i['columns']))}]" if i["columns"] else "")
                                for i in tb["indexes"]
                            ),
                            small,
                        )
                    )
            if tb.get("sample_rows"):
                cols = tb["sample_columns"]
                story.append(Paragraph(f"Sample rows ({len(tb['sample_rows'])})", small))
                data = [[Paragraph("_row", cell), *[Paragraph(esc(c), cell) for c in cols]]]
                for r in tb["sample_rows"]:
                    data.append(
                        [Paragraph(str(r["_row"]), cell), *[Paragraph(esc(_cell(r.get(c))), cell) for c in cols]]
                    )
                idx_w = 12 * mm
                t = LongTable(data, colWidths=[idx_w, *[(W - idx_w) / max(1, len(cols))] * len(cols)], repeatRows=1)
                t.setStyle(grid)
                story.append(t)
    story += [
        Spacer(1, 6 * mm),
        Paragraph(f"Generated by {esc(report['generator'])}. All database access was read-only.", small),
    ]
    doc.build(story)


def render_docx(report: dict[str, Any], output: str | os.PathLike[str]) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10)
    doc.core_properties.title = report["title"]
    doc.core_properties.author = report["generator"]
    doc.add_heading(report["title"], 0)
    doc.add_paragraph("  ·  ".join(f"{k}: {v}" for k, v in _meta_rows(report)))
    if report.get("notes"):
        doc.add_heading("Notes", 1)
        doc.add_paragraph(report["notes"])

    def add_table(headers: list[str], rows: list[list[str]], style_name: str = "Light Grid Accent 1") -> None:
        table = doc.add_table(rows=1, cols=len(headers))
        try:
            table.style = style_name
        except Exception:
            table.style = "Table Grid"
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
        for r in rows:
            cells = table.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
        for row in table.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)

    for entry in report["databases"]:
        info = entry["info"]
        doc.add_heading(info["file_name"], 1)
        add_table(["Property", "Value"], [[k, v] for k, v in _summary_rows(entry)])
        doc.add_heading("Tables", 2)
        add_table(["Table", "Display name", "Columns", "Indexes", "Records", "Description"], _table_rows(entry))
        for tb in entry["tables"]:
            if not (tb.get("columns") or tb.get("sample_rows")):
                continue
            doc.add_heading(tb["display_name"] + (f"  ({tb['name']})" if tb["display_name"] != tb["name"] else ""), 3)
            if tb.get("description"):
                doc.add_paragraph(tb["description"])
            if tb.get("columns"):
                add_table(
                    ["ID", "Column", "Type", "Storage", "Size", "Encoding"],
                    [
                        [c["identifier"], c["name"], c["type"], c["storage"], c["size"] or "", c["encoding"] or ""]
                        for c in tb["columns"]
                    ],
                )
                if tb.get("indexes"):
                    doc.add_paragraph(
                        "Indexes: "
                        + ", ".join(
                            i["name"]
                            + (" (primary)" if i["is_primary"] else "")
                            + (f" [{', '.join(i['columns'])}]" if i["columns"] else "")
                            for i in tb["indexes"]
                        )
                    )
            if tb.get("sample_rows"):
                cols = tb["sample_columns"]
                doc.add_paragraph(f"Sample rows ({len(tb['sample_rows'])})")
                add_table(["_row", *cols], [[r["_row"], *[_cell(r.get(c)) for c in cols]] for r in tb["sample_rows"]])
    doc.add_paragraph(f"Generated by {report['generator']}. All database access was read-only.")
    doc.save(str(output))


def render_xlsx(report: dict[str, Any], output: str | os.PathLike[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from edb_explorer.core.formats import safe_sheet_name

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(name: str, headers: list[str], rows: list[list[Any]]) -> None:
        ws = wb.create_sheet(safe_sheet_name(name, wb))
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        for i, h in enumerate(headers, start=1):
            width = max(len(str(h)), *(len(str(r[i - 1])) for r in rows)) if rows else len(h)
            ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        ws.freeze_panes = "A2"

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Report", report["title"]])
    for k, v in _meta_rows(report):
        summary.append([k, v])
    if report.get("notes"):
        summary.append(["Notes", report["notes"]])
    summary.append([])
    for entry in report["databases"]:
        summary.append([entry["info"]["file_name"]])
        summary.cell(row=summary.max_row, column=1).font = bold
        for k, v in _summary_rows(entry):
            summary.append([k, v])
        summary.append([])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 100

    tables_rows = []
    columns_rows = []
    samples_rows = []
    for entry in report["databases"]:
        fname = entry["info"]["file_name"]
        for t in entry["tables"]:
            tables_rows.append(
                [
                    fname,
                    t["name"],
                    t["display_name"],
                    t["column_count"],
                    t["index_count"],
                    t.get("record_count"),
                    t.get("description") or "",
                ]
            )
            for c in t.get("columns") or []:
                columns_rows.append(
                    [fname, t["name"], c["identifier"], c["name"], c["type"], c["storage"], c["size"], c["encoding"]]
                )
            for r in t.get("sample_rows") or []:
                samples_rows.append(
                    [fname, t["name"], r["_row"], "; ".join(f"{k}={_cell(v)}" for k, v in r.items() if k != "_row")]
                )
    sheet("Tables", ["Database", "Table", "Display name", "Columns", "Indexes", "Records", "Description"], tables_rows)
    if columns_rows:
        sheet("Columns", ["Database", "Table", "ID", "Column", "Type", "Storage", "Size", "Encoding"], columns_rows)
    if samples_rows:
        sheet("Samples", ["Database", "Table", "Row", "Values"], samples_rows)
    wb.save(str(output))


def write_report(report: dict[str, Any], output: str | os.PathLike[str], fmt: str | None = None) -> Path:
    """Render ``report`` into ``output`` (format inferred from the extension when ``fmt`` is None)."""
    out = Path(output)
    fmt = (fmt or out.suffix.lstrip(".") or "html").lower()
    if fmt not in REPORT_FORMATS:
        raise ExportError(f"Unsupported report format {fmt!r} (choose from {', '.join(REPORT_FORMATS)})")
    if out.suffix.lower().lstrip(".") != fmt:
        out = out.with_suffix(f".{fmt}")
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        if fmt == "json":
            out.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
        elif fmt == "md":
            out.write_text(render_markdown(report), encoding="utf-8")
        elif fmt == "txt":
            out.write_text(render_text(report), encoding="utf-8")
        elif fmt == "html":
            out.write_text(render_html(report), encoding="utf-8")
        elif fmt == "pdf":
            render_pdf(report, out)
        elif fmt == "docx":
            render_docx(report, out)
        elif fmt == "xlsx":
            render_xlsx(report, out)
    except OSError as exc:
        raise ExportError(f"Cannot write {out}: {exc}") from exc
    return out


def generate_report(
    databases: list[EdbDatabase],
    output: str | os.PathLike[str],
    fmt: str | None = None,
    options: ReportOptions | None = None,
    progress: ProgressCallback | None = None,
) -> Path:
    """Build and write a report in one call."""
    return write_report(build_report(databases, options, progress), output, fmt)
