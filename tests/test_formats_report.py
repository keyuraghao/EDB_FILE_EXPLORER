from __future__ import annotations

import json
from pathlib import Path

import pytest

from edb_explorer.core import Session
from edb_explorer.core.exceptions import ExportError
from edb_explorer.core.export import EXPORT_FORMATS, export_database, export_rows, export_table
from edb_explorer.core.formats import make_writer, safe_sheet_name
from edb_explorer.core.report import REPORT_FORMATS, ReportOptions, build_report, generate_report, write_report


@pytest.mark.parametrize("fmt", EXPORT_FORMATS)
def test_export_table_every_format(session: Session, tmp_path: Path, fmt: str) -> None:
    db = session.get("srudb")
    out = tmp_path / f"net.{fmt}"
    assert export_table(db, "Network Data Usage", out, fmt, limit=7) == 7
    assert out.exists() and out.stat().st_size > 0
    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(out, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][:3] == ("_row", "AutoIncId", "TimeStamp") and len(rows) == 8
        assert str(rows[1][2]).startswith("2021-01-24")
    elif fmt == "pdf":
        assert out.read_bytes().startswith(b"%PDF")
    elif fmt == "txt":
        text = out.read_text(encoding="utf-8")
        assert "_row | AutoIncId" in text and "7 row(s)" in text
    elif fmt == "json":
        assert len(json.loads(out.read_text(encoding="utf-8"))) == 7


def test_export_rows_selection(session: Session, tmp_path: Path) -> None:
    db = session.get("srudb")
    rows = [r for r in db.fetch("SruDbIdMapTable", limit=10, bytes_mode="raw").rows if r["_row"] in (1, 2)]
    out = tmp_path / "sel.xlsx"
    assert export_rows(rows, ["IdIndex", "IdBlob"], out, "xlsx", db.column_types("SruDbIdMapTable"), "sel") == 2
    from openpyxl import load_workbook

    ws = load_workbook(out, read_only=True).active
    data = list(ws.iter_rows(values_only=True))
    assert data[1] == (1, 2, "S-1-5-18")
    assert export_rows(rows, ["IdIndex", "IdBlob"], tmp_path / "sel.pdf", "pdf") == 2
    with pytest.raises(ValueError, match="Unsupported export format"):
        make_writer("docx", tmp_path / "x.docx", ["a"])


def test_export_database_single_workbook(session: Session, tmp_path: Path) -> None:
    db = session.get("srudb")
    results = export_database(db, tmp_path / "wb", "xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(tmp_path / "wb" / "SRUDB.xlsx", read_only=True)
    assert set(results) == {"SruDbIdMapTable", "SruDbCheckpointTable", "{973F5D5C-1D90-4944-BE8E-24B94231A174}"}
    assert "Network Data Usage" in wb.sheetnames and "SRUM ID Map (apps _ users)" in wb.sheetnames


def test_safe_sheet_name() -> None:
    assert safe_sheet_name("a/b:c*d?e[f]") == "a_b_c_d_e_f_"
    assert len(safe_sheet_name("x" * 50)) == 31


def test_build_report_structure(session: Session) -> None:
    db = session.get("srudb")
    rep = build_report([db], ReportOptions(case_id="C1", analyst="A", sample_rows=2, compute_hash=True))
    assert rep["case_id"] == "C1" and len(rep["databases"]) == 1
    entry = rep["databases"][0]
    assert entry["profile"]["id"] == "srum" and entry["info"]["sha256"]
    assert entry["total_records"] == 29
    names = {t["name"]: t for t in entry["tables"]}
    assert "MSysObjects" not in names
    net = names["{973F5D5C-1D90-4944-BE8E-24B94231A174}"]
    assert net["record_count"] == 25 and len(net["sample_rows"]) == 2 and "TimeStamp" in net["sample_columns"]
    assert net["columns"][1]["type"] == "DateTime"
    lean = build_report(
        [db], ReportOptions(include_schema=False, sample_rows=0, count_records=False, compute_hash=False)
    )
    t = lean["databases"][0]["tables"][0]
    assert "columns" not in t and "sample_rows" not in t


@pytest.mark.parametrize("fmt", REPORT_FORMATS)
def test_write_report_every_format(session: Session, tmp_path: Path, fmt: str) -> None:
    db = session.get("srudb")
    rep = build_report([db], ReportOptions(sample_rows=1, compute_hash=False, notes="hello <notes>"))
    out = write_report(rep, tmp_path / "report", fmt)
    assert out.suffix == f".{fmt}" and out.stat().st_size > 0
    if fmt in ("html", "md", "txt"):
        text = out.read_text(encoding="utf-8")
        assert "SRUDB.dat" in text and "Network Data Usage" in text
        if fmt == "html":
            assert "&lt;notes&gt;" in text
    elif fmt == "json":
        assert json.loads(out.read_text(encoding="utf-8"))["databases"][0]["info"]["file_name"] == "SRUDB.dat"
    elif fmt == "pdf":
        assert out.read_bytes().startswith(b"%PDF")
    elif fmt == "docx":
        from docx import Document

        doc = Document(str(out))
        assert any("SRUDB.dat" in p.text for p in doc.paragraphs)
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        assert {"Summary", "Tables", "Columns", "Samples"} <= set(load_workbook(out, read_only=True).sheetnames)


def test_generate_report_infers_format_and_cancels(session: Session, tmp_path: Path) -> None:
    db = session.get("srudb")
    path = generate_report([db], tmp_path / "r.md", options=ReportOptions(compute_hash=False))
    assert path.suffix == ".md"
    with pytest.raises(ExportError):
        write_report({}, tmp_path / "r.zip", "zip")
    with pytest.raises(ExportError, match="cancelled"):
        generate_report([db], tmp_path / "r.txt", progress=lambda _m: False)
